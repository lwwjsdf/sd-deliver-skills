"""Tests for validate_performance_test_plan.py."""
import openpyxl
import pytest

from design_performance_test import generate_excel_report, generate_word_report
from validate_performance_test_plan import validate_excel, validate_word


def test_validate_word_valid(tmp_path):
    out = tmp_path / "perf.docx"
    generate_word_report(
        output_path=str(out),
        dau=1_000_000,
        daily_events=5_000_000,
        retention_days=365,
        cloud="AWS",
        region="ap-southeast-1",
        include_cdp=True,
        include_ma=True,
    )
    assert validate_word(str(out)) == 0


def test_validate_word_missing_sections(tmp_path):
    from docx import Document

    out = tmp_path / "perf.docx"
    doc = Document()
    doc.add_paragraph("Only introduction")
    doc.save(str(out))
    assert validate_word(str(out)) == 1


def test_validate_excel_valid(tmp_path):
    out = tmp_path / "perf.xlsx"
    generate_excel_report(
        output_path=str(out),
        dau=1_000_000,
        daily_events=5_000_000,
        retention_days=365,
        cloud="AWS",
        region="ap-southeast-1",
        include_cdp=True,
        include_ma=True,
    )
    assert validate_excel(str(out)) == 0


def test_validate_excel_missing_scope(tmp_path):
    out = tmp_path / "perf.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Environment"
    wb.save(out)
    assert validate_excel(str(out)) == 1


def test_validate_excel_empty_scope(tmp_path):
    out = tmp_path / "perf.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scope"
    ws.append(["No.", "Module", "Scenario", "Data Preparation", "Test Steps", "Expected Metrics"])
    wb.save(out)
    assert validate_excel(str(out)) == 1
